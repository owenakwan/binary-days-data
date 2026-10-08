"""
Which ETFs hold each stock (for the "ETFs" tab on the Binary Days stock page).

Downloads the free daily holdings files that State Street (SPDR funds) and ARK publish,
then flips them around: for every ticker, the funds that hold it and its weight in each.
Writes out/etf_holdings.json. Runs once a day (later runs the same day keep the existing file).

iShares, Vanguard and Invesco block automated downloads, so funds like IBB, IWM, VTI and QQQ aren't covered.
"""
import csv, io, json, os, re, sys, time
from datetime import datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

OUT_DIR = os.environ.get("OUT_DIR", "out")
OUT = os.path.join(OUT_DIR, "etf_holdings.json")
CT = ZoneInfo("America/Chicago")
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36"}

SPDR = {
    "SPY": "S&P 500", "SPMD": "S&P MidCap 400", "SPSM": "S&P SmallCap 600",
    "XBI": "S&P Biotech", "XPH": "S&P Pharmaceuticals", "XHE": "S&P Health Care Equipment", "XHS": "S&P Health Care Services",
    "XLV": "Health Care Select Sector", "XLK": "Technology Select Sector", "XLF": "Financial Select Sector", "XLE": "Energy Select Sector",
    "XLI": "Industrial Select Sector", "XLY": "Consumer Discretionary Select Sector", "XLP": "Consumer Staples Select Sector",
    "XLU": "Utilities Select Sector", "XLB": "Materials Select Sector", "XLRE": "Real Estate Select Sector", "XLC": "Communication Services Select Sector",
    "KRE": "S&P Regional Banking", "KBE": "S&P Bank", "XSD": "S&P Semiconductor", "XSW": "S&P Software & Services", "XRT": "S&P Retail",
    "XME": "S&P Metals & Mining", "XOP": "S&P Oil & Gas Exploration", "XHB": "S&P Homebuilders", "XAR": "S&P Aerospace & Defense", "XTN": "S&P Transportation",
}
ARK = {
    "ARKK": ("ARK Innovation", "ARK_INNOVATION_ETF_ARKK_HOLDINGS.csv"),
    "ARKG": ("ARK Genomic Revolution", "ARK_GENOMIC_REVOLUTION_ETF_ARKG_HOLDINGS.csv"),
    "ARKQ": ("ARK Autonomous Tech & Robotics", "ARK_AUTONOMOUS_TECH._&_ROBOTICS_ETF_ARKQ_HOLDINGS.csv"),
    "ARKW": ("ARK Next Generation Internet", "ARK_NEXT_GENERATION_INTERNET_ETF_ARKW_HOLDINGS.csv"),
    "ARKF": ("ARK Blockchain & Fintech", "ARK_BLOCKCHAIN_&_FINTECH_INNOVATION_ETF_ARKF_HOLDINGS.csv"),
    "ARKX": ("ARK Space & Defense", "ARK_SPACE_&_DEFENSE_INNOVATION_ETF_ARKX_HOLDINGS.csv"),
}
TICKER = re.compile(r"^[A-Z]{1,5}([.\-][A-Z])?$")


def norm(t):
    t = str(t or "").strip().upper().split(" ")[0]
    return t.replace("/", ".").replace("-", ".")


def spdr(etf):
    url = f"https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/holdings-daily-us-en-{etf.lower()}.xlsx"
    r = requests.get(url, headers=UA, timeout=60)
    r.raise_for_status()
    ws = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    asof, head, out = "", None, []
    for row in rows:
        cells = [str(c).strip() if c is not None else "" for c in row]
        if not head:
            if cells and cells[0].lower().startswith("holdings") and len(cells) > 1:
                asof = cells[1].replace("As of ", "")
            low = [c.lower() for c in cells]
            if "ticker" in low and "weight" in low:
                head = (low.index("ticker"), low.index("weight"))
            continue
        if not any(cells):
            break
        t, w = norm(cells[head[0]]), cells[head[1]]
        try:
            w = float(w)
        except ValueError:
            continue
        if TICKER.match(t):
            out.append((t, w))
    return asof, out


def ark(etf, file):
    r = requests.get("https://assets.ark-funds.com/fund-documents/funds-etf-csv/" + file, headers=UA, timeout=60)
    r.raise_for_status()
    rows = list(csv.DictReader(io.StringIO(r.text)))
    out, asof = [], ""
    for row in rows:
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        t = norm(row.get("ticker"))
        w = row.get("weight (%)", "").replace("%", "")
        asof = asof or row.get("date", "")
        try:
            w = float(w)
        except ValueError:
            continue
        if TICKER.match(t):
            out.append((t, w))
    return asof, out


def main():
    force = "--force" in sys.argv
    today = datetime.now(CT).strftime("%Y-%m-%d")
    if not force and os.path.exists(OUT):
        try:
            if json.load(open(OUT, encoding="utf-8")).get("day") == today:
                print("ETF holdings already built today; keeping them.")
                return
        except Exception:  # noqa: BLE001
            pass
    funds, by = {}, {}
    jobs = [(e, n, lambda e=e: spdr(e)) for e, n in SPDR.items()] + [(e, n, lambda e=e, f=f: ark(e, f)) for e, (n, f) in ARK.items()]
    for etf, name, fetch in jobs:
        try:
            asof, rows = fetch()
        except Exception as ex:  # noqa: BLE001
            print("skip", etf, ex)
            continue
        if not rows:
            print("skip", etf, "no rows")
            continue
        funds[etf] = {"name": name, "asof": asof, "n": len(rows)}
        for t, w in rows:
            by.setdefault(t, []).append([etf, round(w, 3)])
        print(etf, len(rows), asof)
        time.sleep(0.5)
    if len(funds) < 5:
        print("Too few funds downloaded; keeping the previous file.")
        return
    for t in by:
        by[t].sort(key=lambda x: -x[1])
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump({"day": today, "updated": datetime.now(CT).isoformat(timespec="minutes"), "funds": funds, "by": by}, f, separators=(",", ":"))
    print(f"Wrote {OUT}: {len(funds)} funds, {len(by)} tickers")


if __name__ == "__main__":
    main()
