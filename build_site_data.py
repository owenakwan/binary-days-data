#!/usr/bin/env python3
"""
Builds the website's data files (run by .github/workflows/update-site-data.yml):

  site/data/catalysts.json  - your spreadsheet's open catalysts, plus PDUFA / FDA advisory
                              committee dates discovered in news headlines (marked unverified)
  site/data/earnings.json   - every US earnings report for the next 5 weeks (Finnhub), with each
                              company's recent report times so the site can show its usual time

Nothing secret is written. The Finnhub key stays in GitHub Secrets.

Usage:  python build_site_data.py
Env:    FINNHUB_KEY (earnings + advisory committees), SEC_CONTACT (SEC User-Agent contact), XLSX_PATH, OUT_DIR
"""
import datetime as dt
import email.utils
import html
import json
import os
import re
import sys
import time
import urllib.parse
import xml.etree.ElementTree as ET
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook

CT = ZoneInfo("America/Chicago")
HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.environ.get("XLSX_PATH") or os.path.join(HERE, "biotech_catalysts_master.xlsx")
OUT = os.environ.get("OUT_DIR") or os.path.join(HERE, "out")
FINNHUB_KEY = os.environ.get("FINNHUB_KEY", "").strip()
SEC_UA = f"binary-days data builder {os.environ.get('SEC_CONTACT', '').encode('ascii', 'ignore').decode().strip() or 'contact via github.com/owenakwan'}"
BROWSER_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
WEEKS_AHEAD = 5
HISTORY_DAYS = 400

SESSION = requests.Session()
SESSION.headers.update({"User-Agent": BROWSER_UA})


def log(msg):
    print(msg, flush=True)


# ----------------------------------------------------------------------------- helpers
def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
        except (OverflowError, ValueError):
            return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%y", "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def iso(d):
    return d.isoformat() if d else None


def finnhub(path, **params):
    if not FINNHUB_KEY:
        return None
    params["token"] = FINNHUB_KEY
    for attempt in range(3):
        try:
            r = SESSION.get(f"https://finnhub.io/api/v1/{path}", params=params, timeout=30)
            if r.status_code == 429:
                time.sleep(3 + attempt * 3)
                continue
            if r.status_code in (401, 403):
                log(f"[warn] Finnhub {path}: HTTP {r.status_code} (not on the free plan or bad key)")
                return None
            r.raise_for_status()
            return r.json()
        except Exception as e:  # noqa: BLE001
            log(f"[warn] Finnhub {path}: {e}")
            time.sleep(2)
    return None


def sec_tickers():
    """{TICKER: company title} and a list of (normalized name, ticker) for headline matching."""
    try:
        r = requests.get("https://www.sec.gov/files/company_tickers.json", headers={"User-Agent": SEC_UA}, timeout=30)
        r.raise_for_status()
        rows = list(r.json().values())
    except Exception as e:  # noqa: BLE001
        log(f"[warn] SEC ticker list unavailable: {e}")
        return {}, []
    names, by_name = {}, []
    for v in rows:
        t, title = v["ticker"].upper(), v["title"].strip()
        if title.isupper():
            title = title.title()
        names.setdefault(t, title)
        n = norm_company(title)
        if len(n) >= 6:
            by_name.append((n, t))
    by_name.sort(key=lambda x: -len(x[0]))  # longest names first so "Summit Therapeutics" beats "Summit"
    return names, by_name


SUFFIX = re.compile(r"\b(inc|incorporated|corp|corporation|co|company|ltd|limited|plc|ag|sa|nv|se|holdings?|group|the|n\.v|s\.a)\b\.?", re.I)


def norm_company(s):
    s = SUFFIX.sub(" ", s.lower())
    return re.sub(r"[^a-z0-9 ]", " ", re.sub(r"\s+", " ", s)).strip()


MONTH_NAMES = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
MON3 = {m: i for i, m in enumerate(MONTH_NAMES, 1)}
MON_RE = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"


def month_end(y, m):
    return dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)


def infer_window(text, start, end, today):
    """Read guidance like 'By end of Oct', 'Q4 2026', '2H 2026', 'Early Nov', 'Dec 2026' or 'Any day' as a
    date window (same rules as market_brief.py), so 'could report any day' catalysts show as window open."""
    if end or not text:
        return start, end
    t = str(text).strip().lower().replace("–", "-")
    yr = re.search(r"\b(20\d\d)\b", t)
    y = int(yr.group(1)) if yr else today.year
    def year_for(mo):
        return y if yr or month_end(y, mo) >= today - dt.timedelta(days=60) else y + 1
    if re.search(r"\bany ?day\b|\boverdue\b|\bimminent\b", t):
        return today - dt.timedelta(days=30), today + dt.timedelta(days=30)
    m = re.search(r"\bq([1-4])\b", t)
    if m:
        q = int(m.group(1))
        s, e = dt.date(y, 3 * q - 2, 1), month_end(y, 3 * q)
        if "late" in t: s = s + dt.timedelta(days=45)
        elif "early" in t: e = s + dt.timedelta(days=30)
        elif "mid" in t: s, e = s + dt.timedelta(days=25), s + dt.timedelta(days=65)
        return s, e
    m = re.search(r"\b([12])h\b|\bh([12])\b|\b(first|second) half\b", t)
    if m:
        h = 1 if (m.group(1) or m.group(2) or "") == "1" or m.group(3) == "first" else 2
        return (dt.date(y, 1, 1), dt.date(y, 6, 30)) if h == 1 else (dt.date(y, 7, 1), dt.date(y, 12, 31))
    if re.search(r"year[- ]?end|end of (the )?year", t):
        return max(dt.date(y, 7, 1), today - dt.timedelta(days=60)), dt.date(y, 12, 31)
    m = re.search(r"\b(early|mid|late)[- ]?" + MON_RE, t)
    if m:
        mo = MON3[m.group(2)]; yy = year_for(mo)
        lo, hi = {"early": (1, 10), "mid": (11, 20), "late": (21, 31)}[m.group(1)]
        return dt.date(yy, mo, lo), (dt.date(yy, mo, hi) if hi < 28 else month_end(yy, mo))
    m = re.search(r"\bby (the )?(end of )?" + MON_RE, t)
    if m:
        mo = MON3[m.group(3)]; e = month_end(year_for(mo), mo)
        return e - dt.timedelta(days=60), e
    m = re.fullmatch(MON_RE + r",? (20\d\d)", t)
    if m:
        mo = MON3[m.group(1)]
        return dt.date(y, mo, 1), month_end(y, mo)
    return start, end

# ----------------------------------------------------------------------------- spreadsheet
def load_spreadsheet():
    if not os.path.exists(XLSX):
        log("[warn] spreadsheet not found; catalysts list will only contain discovered items")
        return []
    today = dt.datetime.now(CT).date()
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    ws = wb["Catalysts"] if "Catalysts" in wb.sheetnames else wb.worksheets[0]
    header, out = None, []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            cells = [str(c).strip() if c is not None else "" for c in row]
            if "Ticker" in cells:
                header = cells
            continue
        rec = {header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]}
        tk = str(rec.get("Ticker") or "").strip().upper()
        if not tk or str(rec.get("Status") or "").strip().lower() == "resolved":
            continue
        raw = rec.get("Date")
        date_text = to_date(raw).strftime("%b %d, %Y") if isinstance(raw, (dt.date, dt.datetime)) else str(raw or "").strip()
        start, end = infer_window(date_text, to_date(rec.get("Sort Date")) or to_date(raw), to_date(rec.get("End Date")), today)
        out.append({
            "ticker": tk,
            "company": str(rec.get("Company") or "").strip(),
            "search": str(rec.get("Search Name") or rec.get("Company") or tk).strip(),
            "catalyst": str(rec.get("Catalyst") or "").strip(),
            "date_text": date_text,
            "sort": iso(start),
            "end": iso(end),
            "type": str(rec.get("Type") or "").strip(),
            "status": str(rec.get("Status") or "").strip(),
            "notes": str(rec.get("Notes") or "").strip(),
        })
    wb.close()
    return out


# ----------------------------------------------------------------------------- discovery from news
MONTHS = "January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
FULL_DATE = re.compile(rf"\b({MONTHS})\.?\s+(\d{{1,2}})(?:st|nd|rd|th)?,?\s+(20\d\d)\b", re.I)
EXCH_TICKER = re.compile(r"\((?:NASDAQ|Nasdaq|NYSE|NYSE American|NYSE MKT|Cboe|OTC|OTCQX|OTCQB)[A-Za-z ]*:\s*([A-Z]{1,5})\)")
QUERIES = [
    ('"PDUFA" "target action date"', "PDUFA"),
    ('"PDUFA date"', "PDUFA"),
    ('"PDUFA goal date"', "PDUFA"),
    ('FDA accepts application "PDUFA"', "PDUFA"),
    ('FDA "advisory committee" meeting scheduled', "AdComm"),
]


def parse_when(s):
    if not s:
        return None
    try:
        return email.utils.parsedate_to_datetime(s)
    except (TypeError, ValueError, IndexError):
        return None


def google_news(query):
    url = "https://news.google.com/rss/search?" + urllib.parse.urlencode({"q": query, "hl": "en-US", "gl": "US", "ceid": "US:en"})
    try:
        r = SESSION.get(url, timeout=25)
        r.raise_for_status()
        root = ET.fromstring(r.content)
    except Exception as e:  # noqa: BLE001
        log(f"[warn] Google News {query!r}: {e}")
        return []
    items = []
    for it in root.iter("item"):
        title = html.unescape(it.findtext("title") or "")
        desc = re.sub(r"<[^>]+>", " ", html.unescape(it.findtext("description") or ""))
        items.append({"title": title, "text": title + " " + desc, "link": it.findtext("link") or "",
                      "source": it.findtext("source") or "", "published": parse_when(it.findtext("pubDate"))})
    return items


def parse_full_date(m):
    try:
        return dt.datetime.strptime(f"{m.group(1)[:3].title()} {m.group(2)} {m.group(3)}", "%b %d %Y").date()
    except ValueError:
        return None


def discover(today, names, by_name, known):
    found = {}
    for q, kind in QUERIES:
        for it in google_news(q):
            text = it["text"]
            # take the first future date that appears near the keyword
            kw = re.search(r"PDUFA|target action date|goal date|advisory committee", text, re.I)
            if not kw:
                continue
            cands = [(abs(m.start() - kw.start()), parse_full_date(m)) for m in FULL_DATE.finditer(text)]
            cands = sorted((dist, d) for dist, d in cands if d and today <= d <= today + dt.timedelta(days=400))
            if not cands or cands[0][0] > 160:
                continue
            date = cands[0][1]
            m = EXCH_TICKER.search(text)
            ticker = m.group(1) if m else None
            if not ticker:
                head = norm_company(it["title"])
                for n, t in by_name:
                    if head.startswith(n + " ") or head == n:
                        ticker = t
                        break
            if not ticker:
                continue
            if (ticker, date.isoformat()) in known:
                continue
            key = (ticker, date.isoformat())
            if key in found:
                continue
            title = re.sub(r"\s+-\s+[^-]{2,60}$", "", it["title"])
            found[key] = {
                "ticker": ticker, "company": names.get(ticker, ""), "date": date.isoformat(), "type": kind,
                "headline": title, "link": it["link"], "source": it["source"],
                "published": it["published"].date().isoformat() if it["published"] else None,
            }
        time.sleep(0.8)
    log(f"Discovered {len(found)} dated events from news")
    return sorted(found.values(), key=lambda r: (r["date"], r["ticker"]))


def adcomm_calendar(today, known):
    data = finnhub("fda-advisory-committee-calendar") or []
    out = []
    for r in data if isinstance(data, list) else []:
        d = to_date((r.get("fromDate") or "")[:10])
        if not d or d < today:
            continue
        out.append({"ticker": "", "company": "", "date": d.isoformat(), "type": "AdComm",
                    "headline": (r.get("eventDescription") or "FDA advisory committee meeting").strip()[:240],
                    "link": r.get("url") or "", "source": "FDA (via Finnhub)", "published": None})
    log(f"FDA advisory committee meetings ahead: {len(out)}")
    return out


# ----------------------------------------------------------------------------- company-announced readouts
# "X to Announce Topline Data on Oct. 12", "X to Host Investor Call to Share Phase 3 Results" - usually only 1-3 days' notice.
READOUT_QUERIES = [
    '"to announce" topline data', '"to announce" top-line data', '"to report" topline results', '"to report" top-line',
    '"to host" "investor call" topline', '"to host" "conference call" topline results', '"to host" "investor call" data phase',
    '"to present" topline data', '"to unveil" clinical data', '"to share" topline results', '"to host" "R&D day"',
    '"investor event" "clinical data" to host', '"to report" data phase 3 "webcast"',
]
ANNOUNCE_RE = re.compile(r"\b(to|will)\s+(announce|report|present|share|unveil|host|release|disclose|hold)\b", re.I)
EVENT_RE = re.compile(r"top-?line|clinical data|\bdata\b|results|readout|investor (call|event|day)|r&d day|webcast", re.I)
SKIP_RE = re.compile(r"financial results|quarter(ly)?\b|fiscal|earnings|annual (general |shareholders' )?meeting|inducement|dividend|first patient|enrollment|dosed", re.I)
WEEKDAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
MONTH_DAY = re.compile(rf"\b({MONTHS})\.?\s+(\d{{1,2}})(?:st|nd|rd|th)?\b(?:,?\s+(20\d\d))?", re.I)
TIME_RE = re.compile(r"\b(\d{1,2}(?::\d{2})?\s*(?:a\.?m\.?|p\.?m\.?))\s*(?:\(?\s*(ET|EDT|EST|PT|PDT|PST|CT|CDT|CST)\b)?", re.I)
NEAR_RE = re.compile(r"will host|to host|webcast|conference call|investor call|investor event|will (announce|report|present|release)|to (announce|report|present|release)", re.I)


def dates_in(text, base):
    """Every (position, date) in text. Month-day without a year gets the year that puts it nearest after `base`.
    'today' / 'tomorrow' / weekday names are read relative to `base` (the day the release came out)."""
    out = []
    for m in MONTH_DAY.finditer(text):
        try:
            mo = dt.datetime.strptime(m.group(1)[:3].title(), "%b").month
            day = int(m.group(2))
            if m.group(3):
                d = dt.date(int(m.group(3)), mo, day)
            else:
                d = dt.date(base.year, mo, day)
                if d < base - dt.timedelta(days=3):
                    d = dt.date(base.year + 1, mo, day)
            out.append((m.start(), d))
        except ValueError:
            continue
    low = text.lower()
    for m in re.finditer(r"\btoday\b", low):
        out.append((m.start(), base))
    for m in re.finditer(r"\btomorrow\b", low):
        out.append((m.start(), base + dt.timedelta(days=1)))
    for i, w in enumerate(WEEKDAYS):
        for m in re.finditer(rf"\b{w}\b", low):
            out.append((m.start(), base + dt.timedelta(days=(i - base.weekday()) % 7)))
    return out


def pick_date(text, base, today, horizon=45):
    """The event date: the date mentioned closest to 'will host' / 'webcast' / 'to announce' wording."""
    cands = [(p, d) for p, d in dates_in(text, base) if today <= d <= today + dt.timedelta(days=horizon)]
    if not cands:
        return None, ""
    anchors = [m.start() for m in NEAR_RE.finditer(text)] or [0]
    pos, d = min(cands, key=lambda c: min(abs(c[0] - a) for a in anchors))
    t = TIME_RE.search(text[max(0, pos - 20): pos + 140])
    when = (t.group(1).replace(" ", "").lower() + (" " + t.group(2).upper() if t.group(2) else "")) if t else ""
    return d, when


def page_text(url):
    try:
        r = SESSION.get(url, timeout=20, allow_redirects=True)
        if not r.ok:
            return ""
        txt = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", r.text, flags=re.S | re.I)
        txt = html.unescape(re.sub(r"<[^>]+>", " ", txt))
        return re.sub(r"\s+", " ", txt)
    except Exception:  # noqa: BLE001
        return ""


def release_text(ticker, title, published):
    """Full press-release text for an undated headline: find the same story in Finnhub's company news and open it."""
    start = (published or dt.date.today()) - dt.timedelta(days=4)
    news = finnhub("company-news", symbol=ticker, **{"from": start.isoformat(), "to": dt.date.today().isoformat()}) or []
    want = set(re.findall(r"[a-z0-9]{4,}", title.lower()))
    best, score = None, 0.0
    for n in news if isinstance(news, list) else []:
        have = set(re.findall(r"[a-z0-9]{4,}", (n.get("headline") or "").lower()))
        s = len(want & have) / (len(want) or 1)
        if s > score:
            best, score = n, s
    if not best or score < 0.5:
        return ""
    return (best.get("summary") or "") + " " + page_text(best.get("url") or "")


def ticker_for(text, title, by_name):
    m = EXCH_TICKER.search(text)
    if m:
        return m.group(1)
    head = norm_company(title)
    for n, t in by_name:
        if head.startswith(n + " ") or head == n:
            return t
    return None


def readout_events(today, names, by_name, known):
    found, undated, seen = {}, [], set()
    for q in READOUT_QUERIES:
        for it in google_news(q + " when:10d"):
            title = re.sub(r"\s+-\s+[^-]{2,60}$", "", it["title"])
            if title.lower()[:90] in seen:
                continue
            seen.add(title.lower()[:90])
            if not (ANNOUNCE_RE.search(title) and EVENT_RE.search(title)) or SKIP_RE.search(title):
                continue
            ticker = ticker_for(it["text"], title, by_name)
            if not ticker:
                continue
            pub = it["published"].astimezone(CT).date() if it["published"] else today
            d, when = pick_date(title, pub, today)
            row = {"ticker": ticker, "company": names.get(ticker, ""), "type": "Data / investor event", "headline": title,
                   "link": it["link"], "source": it["source"] or "Google News", "published": pub.isoformat()}
            if d:
                row.update(date=d.isoformat(), time=when)
                found.setdefault((ticker, d.isoformat()), row)
            else:
                undated.append((row, pub))
        time.sleep(0.8)
    # Headlines without a date ("Mirum to Host Investor Call to Share Topline Results"): read the release itself
    looked = set()
    for row, pub in undated[:25]:
        if row["ticker"] in looked or any(k[0] == row["ticker"] for k in found):
            continue
        looked.add(row["ticker"])
        text = release_text(row["ticker"], row["headline"], pub)
        if not text:
            continue
        a = NEAR_RE.search(text)
        d, when = pick_date(text[max(0, a.start() - 400): a.start() + 600] if a else text[:3000], pub, today)
        if d:
            row.update(date=d.isoformat(), time=when)
            found.setdefault((row["ticker"], d.isoformat()), row)
        time.sleep(1.1)
    out = [r for k, r in found.items() if k not in known]
    log(f"Company-announced data events ahead: {len(out)} ({len(undated)} undated headlines, {len(looked)} releases opened)")
    for r in out:
        log(f"   {r['date']} {r.get('time', '')} {r['ticker']}: {r['headline'][:100]}")
    return sorted(out, key=lambda r: (r["date"], r["ticker"]))


# ----------------------------------------------------------------------------- earnings
def monday(d):
    return d - dt.timedelta(days=d.weekday())


def earnings(today, names):
    if not FINNHUB_KEY:
        log("[warn] FINNHUB_KEY missing: earnings calendar skipped")
        return None
    ok_sym = re.compile(r"^[A-Z]{1,5}$")
    start, end = monday(today), monday(today) + dt.timedelta(days=7 * WEEKS_AHEAD - 1)
    rows, seen = [], set()
    d = start
    while d <= end:  # one call per week keeps each response small
        chunk_end = min(d + dt.timedelta(days=6), end)
        res = finnhub("calendar/earnings", **{"from": d.isoformat(), "to": chunk_end.isoformat()}) or {}
        for e in res.get("earningsCalendar", []):
            s = (e.get("symbol") or "").upper()
            if not ok_sym.match(s) or (s, e.get("date")) in seen:
                continue
            seen.add((s, e.get("date")))
            rows.append([s, e.get("date"), e.get("hour") or "", e.get("epsEstimate"), e.get("revenueEstimate"),
                         e.get("quarter"), e.get("year")])
        d = chunk_end + dt.timedelta(days=1)
        time.sleep(1.1)
    log(f"Upcoming earnings rows: {len(rows)}")

    upcoming = {r[0] for r in rows}
    # Report-time history builds up over time: keep what earlier runs saved, add reports whose
    # dates have now passed, and ask Finnhub for history only once a day (it's ~30 calls).
    prev = {}
    try:
        with open(os.path.join(OUT, "earnings.json"), encoding="utf-8") as f:
            prev = json.load(f)
    except (OSError, ValueError):
        pass
    hist = {}
    def add(s, date, hour, act, est):
        if date:
            hist.setdefault(s, {})[date] = [date, hour or "", act, est]
    for s, items in (prev.get("history") or {}).items():
        for x in items:
            add(s, *x[:4])
    for r in prev.get("rows") or []:
        if r[1] and r[1] < today.isoformat() and r[2]:
            add(r[0], r[1], r[2], None, r[3])
    if prev.get("history_fetched") != today.isoformat():
        d = today - dt.timedelta(days=HISTORY_DAYS)
        while d < today:
            chunk_end = min(d + dt.timedelta(days=13), today - dt.timedelta(days=1))
            res = finnhub("calendar/earnings", **{"from": d.isoformat(), "to": chunk_end.isoformat()}) or {}
            for e in res.get("earningsCalendar", []):
                s = (e.get("symbol") or "").upper()
                if s in upcoming:
                    add(s, e.get("date"), e.get("hour"), e.get("epsActual"), e.get("epsEstimate"))
            d = chunk_end + dt.timedelta(days=1)
            time.sleep(1.1)
        fetched = today.isoformat()
    else:
        fetched = prev.get("history_fetched")
    oldest = (today - dt.timedelta(days=450)).isoformat()  # keep about five quarters per company
    history = {s: [x for x in sorted(v.values(), key=lambda x: x[0], reverse=True) if x[0] >= oldest][:6] for s, v in hist.items()}
    history = {s: v for s, v in history.items() if v}
    log(f"Report-time history on file for {len(history)} of {len(upcoming)} companies")

    rows.sort(key=lambda r: (r[1] or "", -(r[4] or 0)))
    return {
        "sample": False,
        "updated": dt.datetime.now(CT).isoformat(timespec="minutes"),
        "from": start.isoformat(), "to": end.isoformat(),
        "fields": ["symbol", "date", "hour", "epsEstimate", "revenueEstimate", "quarter", "year"],
        "rows": rows,
        "history_fields": ["date", "hour", "epsActual", "epsEstimate"],
        "history": history,
        "history_fetched": fetched,
        "names": {s: names[s] for s in upcoming if s in names},
    }


# ----------------------------------------------------------------------------- main
def write(name, obj):
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, separators=(",", ":"), ensure_ascii=False)
    log(f"Wrote {path} ({os.path.getsize(path):,} bytes)")


def main():
    today = dt.datetime.now(CT).date()
    names, by_name = sec_tickers()

    cats = load_spreadsheet()
    if not cats:  # spreadsheet unreachable this run: keep the last good list instead of blanking the board
        try:
            with open(os.path.join(OUT, "catalysts.json"), encoding="utf-8") as f:
                cats = json.load(f).get("catalysts") or []
            log(f"Using the previous catalyst list ({len(cats)} rows)")
        except (OSError, ValueError):
            cats = []
    known = {(c["ticker"], c["sort"]) for c in cats if c["sort"]}
    disc = discover(today, names, by_name, known) + adcomm_calendar(today, known)
    try:
        disc += readout_events(today, names, by_name, known)
    except Exception as e:  # noqa: BLE001  (never let the new scanner break the board)
        log(f"[warn] readout scan failed: {e}")
    disc.sort(key=lambda r: (r["date"], r["ticker"]))
    write("catalysts.json", {
        "sample": False,
        "updated": dt.datetime.now(CT).isoformat(timespec="minutes"),
        "catalysts": cats,
        "discovered": disc,
    })

    earn = earnings(today, names)
    if earn is not None and earn["rows"]:
        write("earnings.json", earn)
    else:
        log("Earnings file left unchanged (no data this run)")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # noqa: BLE001
        log(f"[error] {type(e).__name__}: {e}")
        sys.exit(1)
